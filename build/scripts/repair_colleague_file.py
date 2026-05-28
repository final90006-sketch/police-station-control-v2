"""
repair_colleague_file.py — 修復同仁壞檔（v2.2 第二版：純參照法）

同仁 key 了 92 筆案件，計算欄結構引用崩成 #REF!，統計全 0。
第一版修法（補 calculatedColumnFormula）反而讓 Excel 移除 table，更糟。
第二版改用「純儲存格參照」(J15/D15/B15...) 重寫計算欄：
  - 不需 calculatedColumnFormula（不觸發 table 移除）
  - 純參照刪列自動位移、永不崩 #REF!
  - 其他頁 Tbl案件[案類分類] 結構引用照常運作

本腳本：
  1. 從同仁「原始」壞檔載入（保留 92 筆 input 資料）
  2. 案件 計算欄 R/S/T/U/V/X 重寫為純參照
  3. 毒調 計算欄 K/L/M/N 重寫為純參照
  4. table 不動（保持原本有效結構）
  5. 另存「_已修復」版
"""
import sys
import os
import re
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent))

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter, column_index_from_string

from phases._base import load_config

SRC = r"D:\複本 PoliceStation_v2.2_部署版.xlsx"
DST = r"D:\複本 PoliceStation_v2.2_已修復_正式版.xlsx"


def structured_to_plain(formula, row, name_to_col):
    def repl(m):
        col = name_to_col.get(m.group(1))
        return f"{col}{row}" if col else m.group(0)
    return re.sub(r"\[@([^\]]+)\]", repl, formula)


def repair_table_cells(ws, tbl_name, calc_formulas_struct, log_prefix):
    """重寫計算欄 cell 為純參照（不動 table 定義）。

    calc_formulas_struct: {欄名: 含 [@欄名] 的公式（不含=）}
    """
    tbl = ws.tables[tbl_name]
    m = re.match(r"([A-Z]+)(\d+):([A-Z]+)(\d+)", tbl.ref)
    c1, r1, c2, r2 = m.group(1), int(m.group(2)), m.group(3), int(m.group(4))
    header_row, data_start, data_end = r1, r1 + 1, r2

    # 欄名 → 欄字母
    col_start = column_index_from_string(c1)
    col_end = column_index_from_string(c2)
    name_to_col = {}
    for ci in range(col_start, col_end + 1):
        cl = get_column_letter(ci)
        hdr = ws[f"{cl}{header_row}"].value
        if hdr:
            name_to_col[hdr] = cl

    rewritten = 0
    for name, fstruct in calc_formulas_struct.items():
        if name not in name_to_col:
            print(f"  ⚠ {log_prefix} 找不到欄「{name}」")
            continue
        col = name_to_col[name]
        for r in range(data_start, data_end + 1):
            ws[f"{col}{r}"] = "=" + structured_to_plain(fstruct, r, name_to_col)
            rewritten += 1
    print(f"  {log_prefix} {tbl_name}: 重寫 {rewritten} 計算欄 cell（純參照，table 定義不動）")


def main():
    print("===== 修復同仁壞檔（純參照法）開始 =====")
    if not os.path.exists(SRC):
        print(f"  ✗ 來源不存在：{SRC}")
        return

    # 案件 計算欄公式（從 config，含 [@欄名]）
    cfg = load_config("phase2_data_tables")
    case_calc = {c["name"]: c["formula"].lstrip("=")
                 for c in cfg["page5_case_database"]["columns"]
                 if c["kind"] == "calc" and c.get("formula")}

    # 毒調 計算欄公式（含 [@欄名]）
    drug_calc = {
        "距今天數": 'IFERROR(IF([@最後到驗日]="","",今日-[@最後到驗日]),"")',
        "是否到驗": ('IFERROR(IF([@管制情形]="",0,IF(OR(SUMPRODUCT(--ISNUMBER(SEARCH('
                    '{"已驗","已到驗","通緝","強採","在監"},[@管制情形])))>0,'
                    'AND(ISNUMBER([@距今天數]),[@距今天數]<=30)),1,0)),0)'),
        "狀態燈": ('IFERROR(IF([@管制情形]="","○ 未填",'
                  'IF(ISNUMBER(SEARCH("解除",[@管制情形])),"○ 結束",'
                  'IF(AND(ISNUMBER([@距今天數]),[@距今天數]>30),"● 紅",'
                  'IF([@是否到驗]=1,"● 綠","⚠ 黃")))),"")'),
        "候選旗標": ('IFERROR(IF(OR(ISNUMBER(SEARCH("未到驗",[@管制情形])),'
                    'AND(ISNUMBER([@距今天數]),[@距今天數]>30)),1,0),0)'),
    }

    wb = load_workbook(SRC)
    repair_table_cells(wb["案件資料庫"], "Tbl案件", case_calc, "[案件]")
    if "毒品調驗人口管制" in wb.sheetnames:
        ws_d = wb["毒品調驗人口管制"]
        if "Tbl毒調" in ws_d.tables:
            repair_table_cells(ws_d, "Tbl毒調", drug_calc, "[毒調]")

    wb.calculation.fullCalcOnLoad = True
    wb.save(DST)

    # 驗證
    wb2 = load_workbook(DST)
    ws2 = wb2["案件資料庫"]
    import zipfile
    with zipfile.ZipFile(DST) as z:
        t1 = z.read("xl/tables/table1.xml").decode("utf-8")
    print()
    print("=== 驗證 ===")
    for col, name in [("R","是否破獲"),("S","歸屬年度"),("T","案類分類")]:
        f = str(ws2[f"{col}15"].value or "")
        ok = "[@" not in f and "#REF!" not in f
        print(f"  {col}15 {name}: {'✓ 純參照' if ok else '🔴'} {f[:45]}")
    print(f"  table1.xml calculatedColumnFormula 數={t1.count('calculatedColumnFormula')}（應為 0）")
    # 抽查資料還在
    print(f"  資料抽查: row15 編號={ws2['A15'].value} / row106 編號={ws2['A106'].value}")
    print(f"  ✓ 修復完成：{os.path.basename(DST)} ({os.path.getsize(DST)/1024:.1f} KB)")
    print("===== 完成 =====")


if __name__ == "__main__":
    main()
