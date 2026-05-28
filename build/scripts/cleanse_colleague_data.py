"""
cleanse_colleague_data.py — 同仁資料完整清洗（v2.2）

修 3 大資料問題（讓統計真正算得出來）：
  1. 發生時間/破獲時間 文字 → 真 datetime
     （格式 YYYY/MM/DD HHMM 時間沒冒號、typo 等）
  2. 案類 自由打字 → 下拉字典標準值（原文保留到「自填案類」欄）
  3. 發生管轄 非標準值（法院交辦偵查隊...）→ 本轄/他轄

並重寫計算欄為純參照（防 #REF!）。
產生逐筆變更報告 cleanse_report.txt 供稽核。
"""
import sys
import os
import re
from datetime import datetime
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent))

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter, column_index_from_string

from phases._base import load_config

SRC = r"D:\複本 PoliceStation_v2.2_部署版.xlsx"
DST = r"D:\複本 PoliceStation_v2.2_最終清洗版.xlsx"
REPORT = r"D:\複本_資料清洗報告.txt"

# 案類關鍵字 → 字典標準值（順序＝優先，特定在前）
CASE_MAP_RULES = [
    ("機車竊盜", "竊盜(機車)"),
    ("竊盜",     "竊盜(普通)"),
    ("詐欺",     "詐欺(普通)"),
    ("毒品",     "毒品危害防制條例"),
    ("公共危險", "公共危險(危險駕駛)"),
    ("肇事逃逸", "肇事逃逸"),
    ("駕駛過失", "駕駛過失(含致傷重傷死亡)"),
    ("家庭暴力", "一般傷害(不含駕駛過失)"),
    ("傷害",     "一般傷害(不含駕駛過失)"),
    ("強制",     "強制罪"),
    ("妨害自由", "妨害自由"),
    ("賭博",     "賭博"),
    ("侵占",     "侵占"),
    ("妨害風化", "妨害風化"),
    ("毀棄損壞", "毀棄損壞"),
    ("毀損",     "毀棄損壞"),
    ("保護令",   "違反保護令罪"),
    ("背信",     "背信"),
    ("性騷",     "性騷擾防治法"),
    ("名譽",     "妨害名譽(公然侮辱、誹謗)"),
    ("電腦使用", "妨害電腦使用"),
    ("偽造文書", "偽造文書"),
    ("違造文書", "偽造文書"),
    ("妨害秩序", "妨害秩序"),
    ("恐嚇",     "恐嚇取財(一般)"),
]


def map_case_type(raw, valid_names):
    """自由打字 → 字典標準值。找不到 → 其他(自填)"""
    if raw is None or str(raw).strip() == "":
        return None
    s = str(raw).strip()
    # 已經是標準值就不動
    if s in valid_names:
        return s
    for kw, std in CASE_MAP_RULES:
        if kw in s:
            return std
    return "其他(自填)"


def parse_date(s):
    """文字日期 → datetime。回傳 (datetime|None, 是否原本就是日期)"""
    if isinstance(s, datetime):
        return s, True
    if s is None or str(s).strip() == "":
        return None, True
    s = str(s).strip()
    for fmt in ("%Y-%m-%d %H:%M:%S", "%Y-%m-%d %H:%M",
                "%Y/%m/%d %H:%M", "%Y-%m-%d", "%Y/%m/%d"):
        try:
            return datetime.strptime(s, fmt), False
        except ValueError:
            pass

    def _norm_year(y):
        """民國年 → 西元（< 1911 視為民國，+1911）"""
        return y + 1911 if y < 1911 else y

    # YYYY/MM/DD HHMM（時間無冒號，3-4 位）含民國年
    m = re.match(r"(\d{2,4})[/\-.](\d{1,2})[/\-.](\d{1,2})\s+(\d{3,4})$", s)
    if m:
        y, mo, d, t = _norm_year(int(m[1])), int(m[2]), int(m[3]), m[4].zfill(4)
        hh, mm = int(t[:2]), int(t[2:])
        if hh < 24 and mm < 60:
            try:
                return datetime(y, mo, d, hh, mm), False
            except ValueError:
                pass
        try:
            return datetime(y, mo, d), False
        except ValueError:
            pass
    # YYYY/MM/DD + 任何尾巴（含 5 位 typo）→ 取日期，含民國年
    m = re.match(r"(\d{2,4})[/\-.](\d{1,2})[/\-.](\d{1,2})", s)
    if m:
        try:
            return datetime(_norm_year(int(m[1])), int(m[2]), int(m[3])), False
        except ValueError:
            pass
    return None, False   # 完全無法解析


def structured_to_plain(formula, row, name_to_col):
    def repl(m):
        col = name_to_col.get(m.group(1))
        return f"{col}{row}" if col else m.group(0)
    return re.sub(r"\[@([^\]]+)\]", repl, formula)


def main():
    print("===== 同仁資料完整清洗 開始 =====")
    if not os.path.exists(SRC):
        print(f"  ✗ 來源不存在：{SRC}")
        return

    cfg_dict = load_config("phase1_5_settings")
    valid_case_names = {ct[1] for ct in cfg_dict["section_d"]["case_types"]}

    cfg2 = load_config("phase2_data_tables")
    case_calc = {c["name"]: c["formula"].lstrip("=")
                 for c in cfg2["page5_case_database"]["columns"]
                 if c["kind"] == "calc" and c.get("formula")}

    wb = load_workbook(SRC)
    ws = wb["案件資料庫"]
    tbl = ws.tables["Tbl案件"]
    m = re.match(r"([A-Z]+)(\d+):([A-Z]+)(\d+)", tbl.ref)
    c1, r1, c2, r2 = m.group(1), int(m.group(2)), m.group(3), int(m.group(4))
    header_row, data_start, data_end = r1, r1 + 1, r2

    # 欄名 → 欄字母
    col_start = column_index_from_string(c1)
    col_end = column_index_from_string(c2)
    name_to_col = {}
    for ci in range(col_start, col_end + 1):
        cl = get_column_letter(ci)
        hdr = ws[f"{cl}{header_row}"].value
        if hdr:
            name_to_col[hdr] = cl

    COL = name_to_col  # alias
    report = []
    case_changes = []
    date_changes = []
    juris_changes = []
    date_fail = []

    for r in range(data_start, data_end + 1):
        bno = ws[f"{COL['編號']}{r}"].value
        if bno is None or str(bno).strip() == "":
            continue   # 空白列跳過

        # --- 1. 案類映射 ---
        raw_case = ws[f"{COL['案類']}{r}"].value
        if raw_case and str(raw_case).strip():
            mapped = map_case_type(raw_case, valid_case_names)
            if mapped and mapped != str(raw_case).strip():
                # 原文保留到 自填案類
                if "自填案類" in COL and not ws[f"{COL['自填案類']}{r}"].value:
                    ws[f"{COL['自填案類']}{r}"] = raw_case
                ws[f"{COL['案類']}{r}"] = mapped
                case_changes.append((r, str(raw_case), mapped))

        # --- 2. 日期解析（發生時間 + 破獲時間）---
        for date_col in ["發生時間", "破獲時間"]:
            if date_col in COL:
                cell = f"{COL[date_col]}{r}"
                v = ws[cell].value
                dt, was_date = parse_date(v)
                if dt is not None and not was_date:
                    ws[cell] = dt
                    ws[cell].number_format = "yyyy/mm/dd h\"時\""
                    date_changes.append((r, date_col, str(v), dt.strftime("%Y/%m/%d %H:%M")))
                elif dt is None and v is not None and str(v).strip():
                    date_fail.append((r, date_col, str(v)))

        # --- 3. 發生管轄 + 查獲管轄 正規化 ---
        # 業務規則（使用者校正 v2）：法院交辦案 = 本轄發生（發生管轄=本轄）。
        #   查獲管轄：雖偵查隊給的，但「我們有寫破獲時間就是我們破獲」→
        #     有破獲時間 → 查獲管轄=本轄；無破獲時間 → 他轄（隊破）。
        jcol = f"{COL['發生管轄']}{r}"
        ccol = f"{COL['查獲管轄']}{r}"
        jv = ws[jcol].value
        if jv and str(jv).strip() not in ("本轄", "他轄"):
            old = str(jv)
            ws[jcol] = "本轄"          # 法院交辦案 = 本轄發生
            new_catch = ""
            if "隊破" in old or "偵查隊" in old:
                bv = ws[f"{COL['破獲時間']}{r}"].value
                has_break = bv is not None and str(bv).strip() != ""
                ws[ccol] = "本轄" if has_break else "他轄"
                new_catch = (" ｜ 查獲管轄→本轄(有破獲時間=我們破)" if has_break
                             else " ｜ 查獲管轄→他轄(無破獲時間=隊破)")
            juris_changes.append((r, old, "本轄" + new_catch))

    # --- 4. 計算欄重寫純參照（防 #REF!）---
    for name, fstruct in case_calc.items():
        if name in COL:
            col = COL[name]
            for r in range(data_start, data_end + 1):
                ws[f"{col}{r}"] = "=" + structured_to_plain(fstruct, r, name_to_col)

    # 毒調計算欄也順手修
    if "毒品調驗人口管制" in wb.sheetnames:
        wsd = wb["毒品調驗人口管制"]
        if "Tbl毒調" in wsd.tables:
            dcalc = {
                "距今天數": 'IFERROR(IF([@最後到驗日]="","",今日-[@最後到驗日]),"")',
                "是否到驗": ('IFERROR(IF([@管制情形]="",0,IF(OR(SUMPRODUCT(--ISNUMBER(SEARCH('
                            '{"已驗","已到驗","通緝","強採","在監"},[@管制情形])))>0,'
                            'AND(ISNUMBER([@距今天數]),[@距今天數]<=30)),1,0)),0)'),
                "狀態燈": ('IFERROR(IF([@管制情形]="","○ 未填",'
                          'IF(ISNUMBER(SEARCH("解除",[@管制情形])),"○ 結束",'
                          'IF(AND(ISNUMBER([@距今天數]),[@距今天數]>30),"● 紅",'
                          'IF([@是否到驗]=1,"● 綠","⚠ 黃")))),"")'),
                "候選旗標": ('IFERROR(IF(OR(ISNUMBER(SEARCH("未到驗",[@管制情形])),'
                            'AND(ISNUMBER([@距今天數]),[@距今天數]>30)),1,0),0)'),
            }
            dm = re.match(r"([A-Z]+)(\d+):([A-Z]+)(\d+)", wsd.tables["Tbl毒調"].ref)
            d_cs, d_ce = column_index_from_string(dm.group(1)), column_index_from_string(dm.group(3))
            d_hdr, d_ds, d_de = int(dm.group(2)), int(dm.group(2))+1, int(dm.group(4))
            dname_col = {}
            for ci in range(d_cs, d_ce+1):
                cl = get_column_letter(ci); h = wsd[f"{cl}{d_hdr}"].value
                if h: dname_col[h] = cl
            for name, fstruct in dcalc.items():
                if name in dname_col:
                    for r in range(d_ds, d_de+1):
                        wsd[f"{dname_col[name]}{r}"] = "=" + structured_to_plain(fstruct, r, dname_col)

    wb.calculation.fullCalcOnLoad = True
    wb.save(DST)

    # === 變更報告 ===
    lines = []
    lines.append("=" * 60)
    lines.append("同仁資料清洗報告")
    lines.append("=" * 60)
    lines.append(f"\n【1. 日期解析】文字 → 真日期 共 {len(date_changes)} 筆")
    for r, c, old, new in date_changes:
        lines.append(f"  row{r} {c}: 「{old}」→ {new}")
    if date_fail:
        lines.append(f"\n  ⚠ 無法解析 {len(date_fail)} 筆（請手動修）：")
        for r, c, old in date_fail:
            lines.append(f"    row{r} {c}: 「{old}」")

    lines.append(f"\n【2. 案類映射】自由打字 → 字典值 共 {len(case_changes)} 筆")
    lines.append("  （原文已保留到「自填案類」欄）")
    seen = {}
    for r, old, new in case_changes:
        seen.setdefault((old, new), []).append(r)
    for (old, new), rows in sorted(seen.items(), key=lambda x: -len(x[1])):
        lines.append(f"  「{old}」→「{new}」 ({len(rows)} 筆: row {rows[0]}...)")

    lines.append(f"\n【3. 發生管轄正規化】非標準值 → 本轄 共 {len(juris_changes)} 筆")
    lines.append("  ⚠⚠ 預設改「本轄」，若實際是他轄請手動改！")
    for r, old, new in juris_changes:
        lines.append(f"  row{r}: 「{old}」→ {new}")

    report_text = "\n".join(lines)
    with open(REPORT, "w", encoding="utf-8") as f:
        f.write(report_text)

    print(report_text)
    print(f"\n  ✓ 清洗完成：{os.path.basename(DST)} ({os.path.getsize(DST)/1024:.1f} KB)")
    print(f"  ✓ 報告：{REPORT}")
    print("===== 完成 =====")


if __name__ == "__main__":
    main()
