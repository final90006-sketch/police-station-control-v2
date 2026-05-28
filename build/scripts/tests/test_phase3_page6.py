"""
test_phase3_page6.py — 頁 6 刑案管制 v2.1 兩區段並呈版 — 6 大類測試
"""
import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

from openpyxl import load_workbook
from phases._base import OUTPUT_PATH
from tests.framework import TestSuite, Category


SHEET = "刑案管制"


def run():
    suite = TestSuite("頁 6 刑案管制（兩區段並呈）")
    wb = load_workbook(OUTPUT_PATH)
    ws = wb[SHEET]

    # =========================================================
    # 1. 公式語法
    # =========================================================
    cat = Category.SYNTAX

    a1 = str(ws["A1"].value or "")
    suite.assert_true("Banner 含「刑案管制」", cat, "刑案管制" in a1)
    suite.assert_true("Banner 含「兩區段並呈」", cat, "兩區段並呈" in a1)

    a2 = str(ws["A2"].value or "")
    for nr in ["警察局名稱", "分局名稱", "派出所名稱"]:
        suite.assert_true(f"sub-banner 引用 '{nr}'", cat, nr in a2)

    # 維護指引（row 4）
    a4 = str(ws["A4"].value or "")
    suite.assert_true("維護指引含「派出所主管兩階段」", cat,
                      "派出所主管兩階段" in a4)

    # 流程圖 row 5（3 方塊）
    flow_cells = {"A5": "尚未偵破", "D5": "已破獲未移送", "G5": "已移送"}
    for cell, stage in flow_cells.items():
        v = str(ws[cell].value or "")
        suite.assert_true(f"流程圖 {cell} 含「{stage}」", cat, stage in v,
                          detail=v[:60])

    # KPI 縮影 row 7（3 卡）
    for cell, label in [("A7", "🔴 尚未偵破"), ("D7", "🟡 已破獲未移送"), ("G7", "✓ 已移送")]:
        v = str(ws[cell].value or "")
        suite.assert_true(f"KPI {cell} 含「{label}」", cat, label in v,
                          detail=v[:60])
        suite.assert_true(f"KPI {cell} 公式存在", cat,
                          v.startswith("="), detail=v[:60])

    # Section 1 title row 9
    a9 = str(ws["A9"].value or "")
    suite.assert_true("Section 1 含「尚未偵破清單」", cat,
                      "尚未偵破清單" in a9, detail=a9[:60])

    # Section 1 header row 10 (9 cols)
    expected_headers = ["編號", "案類", "發生時間", "發生地", "破獲時間",
                        "案件狀況", "偵辦進度", "承辦人", "自填案類"]
    for i, h in enumerate(expected_headers, start=1):
        col = chr(ord("A") + i - 1)
        suite.assert_eq(f"Section 1 header {col}10", cat,
                        ws[f"{col}10"].value, h)

    # Section 1 data rows 11-25（15 列）— 各 INDEX/MATCH from S 欄
    for k in [1, 8, 15]:
        row = 10 + k
        for col in ["A", "B", "F"]:
            f = str(ws[f"{col}{row}"].value or "")
            suite.assert_true(f"Section 1 {col}{row} 用 INDEX/MATCH", cat,
                              "INDEX" in f and "MATCH" in f, detail=f[:60])
            suite.assert_true(f"Section 1 {col}{row} 含 IFERROR", cat,
                              "IFERROR" in f)
            suite.assert_true(f"Section 1 {col}{row} 引用 helper S", cat,
                              "$S$100" in f, detail=f[:80])

    # Section 2 title row 27
    a27 = str(ws["A27"].value or "")
    suite.assert_true("Section 2 含「已破獲未移送清單」", cat,
                      "已破獲未移送清單" in a27, detail=a27[:60])

    # Section 2 header row 28 (重複 9 cols)
    for i, h in enumerate(expected_headers, start=1):
        col = chr(ord("A") + i - 1)
        suite.assert_eq(f"Section 2 header {col}28", cat,
                        ws[f"{col}28"].value, h)

    # Section 2 data rows 29-43 — 各 INDEX/MATCH from U 欄
    for k in [1, 8, 15]:
        row = 28 + k
        for col in ["A", "B", "F"]:
            f = str(ws[f"{col}{row}"].value or "")
            suite.assert_true(f"Section 2 {col}{row} 用 INDEX/MATCH", cat,
                              "INDEX" in f and "MATCH" in f, detail=f[:60])
            suite.assert_true(f"Section 2 {col}{row} 引用 helper U", cat,
                              "$U$100" in f, detail=f[:80])

    # Helper area row 100：R/S/T/U
    r100 = str(ws["R100"].value or "")
    suite.assert_true("Helper R100 用 AND + 尚未偵破", cat,
                      "尚未偵破" in r100 and "AND" in r100, detail=r100[:80])
    t100 = str(ws["T100"].value or "")
    suite.assert_true("Helper T100 用 AND + 已破獲未移送", cat,
                      "已破獲未移送" in t100 and "AND" in t100, detail=t100[:80])

    # =========================================================
    # 2. 邊界資料
    # =========================================================
    cat = Category.BOUNDARY
    # 無下拉了（移除）
    has_view_dv = False
    for dv in ws.data_validations.dataValidation:
        if dv.formula1 and "尚未偵破" in str(dv.formula1) and "已移送" in str(dv.formula1):
            has_view_dv = True
    suite.assert_true("已移除舊視圖下拉（v2.1 改兩區段並呈）", cat, not has_view_dv)

    # KPI 縮影 IFERROR 包覆
    for cell in ["A7", "D7", "G7"]:
        v = str(ws[cell].value or "")
        suite.assert_true(f"KPI {cell} IFERROR 容錯", cat, "IFERROR" in v)

    # =========================================================
    # 3. 性能基準
    # =========================================================
    cat = Category.PERFORMANCE
    fcount = sum(1 for row in ws.iter_rows() for c in row
                 if c.value and isinstance(c.value, str) and c.value.startswith("="))
    suite.assert_le("頁 6 公式總數 ≤ 1500", cat, fcount, 1500,
                    detail=f"actual={fcount}")
    # 兩區段 helper 各 200 列 × 4 公式 ≈ 800 + 主表 30 列 × 9 ≈ 270 + 縮影 3

    has_today = any("TODAY()" in str(c.value or "")
                    for row in ws.iter_rows() for c in row)
    suite.assert_true("不直接呼叫 TODAY()", cat, not has_today)

    # =========================================================
    # 4. 跨平台相容
    # =========================================================
    cat = Category.CROSS_PLATFORM
    has_filter = any("FILTER(" in str(c.value or "") or "SORT(" in str(c.value or "")
                     for row in ws.iter_rows() for c in row)
    suite.assert_true("無 FILTER/SORT 依賴", cat, not has_filter)

    wb_nrs = set(wb.defined_names)
    for nr in ["今日", "警察局名稱"]:
        suite.assert_true(f"命名範圍 '{nr}' 存在", cat, nr in wb_nrs)

    f = str(ws["A11"].value or "")
    suite.assert_true("使用 Tbl案件 結構引用", cat, "Tbl案件[" in f)

    # =========================================================
    # 5. 真實情境
    # =========================================================
    cat = Category.REAL_WORLD
    # 凍結 = None
    suite.assert_eq("凍結窗格 = None", cat, ws.freeze_panes, None)

    # 頁尾
    f = str(ws["A45"].value or "")
    suite.assert_true("頁尾含 KKEVIN-LIN", cat, "KKEVIN-LIN" in f)
    suite.assert_true("頁尾說明 v2.1 兩區段並呈", cat,
                      "v2.1" in f and "兩區段" in f, detail=f[:80])

    # =========================================================
    # 6. 錯誤恢復
    # =========================================================
    cat = Category.ERROR_RECOVERY
    # 兩區段所有主表公式都包 IFERROR
    sample_cells = ["A11", "F11", "I25", "A29", "F29", "I43"]
    for cell in sample_cells:
        v = str(ws[cell].value or "")
        suite.assert_true(f"{cell} IFERROR 容錯", cat, "IFERROR" in v,
                          detail=v[:60])

    # Helper R/S/T/U IFERROR 包覆
    for cell in ["R100", "T100"]:
        v = str(ws[cell].value or "")
        suite.assert_true(f"Helper {cell} IFERROR", cat, "IFERROR" in v)

    return suite


if __name__ == "__main__":
    suite = run()
    print(suite.report())
    sys.exit(0 if suite.is_all_pass() else 1)
