"""test_phase4_page4.py — 頁 4 對上級機關上呈 6 大類測試"""
import sys
from pathlib import Path
sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

from openpyxl import load_workbook
from phases._base import OUTPUT_PATH
from tests.framework import TestSuite, Category


TARGETS = ["對分局", "對警察局", "對警政署"]


def run():
    suite = TestSuite("頁 4 對上級機關上呈")
    wb = load_workbook(OUTPUT_PATH)
    ws = wb["對上級機關上呈"]

    # === 1. 公式語法 ===
    cat = Category.SYNTAX
    a1 = str(ws["A1"].value or "")
    suite.assert_true("Banner 含「對上級機關上呈」", cat, "對上級機關上呈" in a1)

    # B3 下拉預設值
    suite.assert_true("B3 預設值合理", cat, ws["B3"].value in TARGETS)

    # DataValidation B3 = 3 targets
    has_dv = False
    for dv in ws.data_validations.dataValidation:
        f1 = str(dv.formula1 or "")
        if all(t in f1 for t in TARGETS):
            for r in dv.sqref.ranges:
                if "B3" in str(r):
                    has_dv = True
    suite.assert_true("B3 有 3 對象 DataValidation", cat, has_dv)

    # 動態抬頭 C3 用 CHOOSE
    c3 = str(ws["C3"].value or "")
    suite.assert_true("C3 動態抬頭用 CHOOSE", cat, "CHOOSE" in c3 and "B3" in c3)

    # 公文標題 A6 用 CHOOSE
    a6 = str(ws["A6"].value or "")
    suite.assert_true("公文標題 A6 用 CHOOSE", cat,
                      "CHOOSE" in a6 and "派出所名稱" in a6)

    # 受文者 / 主旨 (rows 7/8) 用 CHOOSE
    for r in [7, 8]:
        v = str(ws[f"C{r}"].value or "")
        suite.assert_true(f"row {r} C{r} CHOOSE", cat,
                          "CHOOSE" in v, detail=v[:60])

    # 8 KPI 列 (rows 12-19)
    for r in range(12, 20):
        no = ws[f"A{r}"].value
        name = ws[f"B{r}"].value
        suite.assert_true(f"KPI row {r} 編號+名稱存在", cat,
                          no is not None and name is not None,
                          detail=f"A{r}={no!r} B{r}={name!r}")
        v = str(ws[f"C{r}"].value or "")
        suite.assert_true(f"KPI row {r} 數值公式存在", cat, v.startswith("="),
                          detail=v[:60])

    # 4 行附件 (rows 22-25) 用 CHOOSE
    for r in range(22, 26):
        v = str(ws[f"A{r}"].value or "")
        suite.assert_true(f"附件 row {r} CHOOSE", cat,
                          "CHOOSE" in v, detail=v[:60])

    # === 2. 邊界 ===
    cat = Category.BOUNDARY
    # 8 KPI 全部
    suite.assert_eq("KPI 表 8 列", cat,
                    sum(1 for r in range(12, 20) if ws[f"B{r}"].value), 8)

    # === 3. 性能 ===
    cat = Category.PERFORMANCE
    fcount = sum(1 for row in ws.iter_rows() for c in row
                 if c.value and isinstance(c.value, str) and c.value.startswith("="))
    suite.assert_le("頁 4 公式總數 ≤ 80", cat, fcount, 80, detail=f"actual={fcount}")

    has_today = any("TODAY()" in str(c.value or "")
                    for row in ws.iter_rows() for c in row)
    suite.assert_true("不直接呼叫 TODAY()", cat, not has_today)

    # === 4. 跨平台 ===
    cat = Category.CROSS_PLATFORM
    has_filter = any("FILTER(" in str(c.value or "") or "SORT(" in str(c.value or "")
                     for row in ws.iter_rows() for c in row)
    suite.assert_true("無 FILTER/SORT 依賴", cat, not has_filter)

    # CHOOSE 跨版本通用（Excel 2010+）
    has_choose = any("CHOOSE(" in str(c.value or "")
                     for row in ws.iter_rows() for c in row)
    suite.assert_true("使用 CHOOSE（跨版本通用）", cat, has_choose)

    # === 5. 真實情境 ===
    cat = Category.REAL_WORLD
    # 3 對象都在下拉選項
    for t in TARGETS:
        in_dv = False
        for dv in ws.data_validations.dataValidation:
            if t in str(dv.formula1 or ""):
                in_dv = True
        suite.assert_true(f"對象「{t}」在下拉", cat, in_dv)

    # Section titles
    for r, name in [(5, "公文抬頭"), (10, "主要 KPI"), (21, "必附附件")]:
        v = str(ws[f"A{r}"].value or "")
        suite.assert_true(f"Section row {r} 含「{name}」", cat, name in v,
                          detail=v[:40])

    # 凍結 = None
    suite.assert_eq("凍結窗格 = None", cat, ws.freeze_panes, None)

    # === 6. 錯誤恢復 ===
    cat = Category.ERROR_RECOVERY
    # 達標狀態欄用 IFERROR (對 G12 抽樣)
    g12 = str(ws["G12"].value or "")
    suite.assert_true("G12 達標狀態 IFERROR", cat, "IFERROR" in g12,
                      detail=g12[:60])

    return suite


if __name__ == "__main__":
    suite = run()
    print(suite.report())
    sys.exit(0 if suite.is_all_pass() else 1)
