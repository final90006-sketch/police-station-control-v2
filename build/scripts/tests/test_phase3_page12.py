"""
test_phase3_page12.py — 頁 12 跨期間趨勢 6 大類高強度測試
"""
import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

from openpyxl import load_workbook

from phases._base import OUTPUT_PATH
from tests.framework import TestSuite, Category


SHEET = "跨期間趨勢分析"
PERIOD_CELL = "B3"
PERIODS = ["6 月", "12 月", "24 月"]


def run():
    suite = TestSuite("頁 12 跨期間趨勢")
    wb = load_workbook(OUTPUT_PATH)
    ws = wb[SHEET]

    # =========================================================
    # 1. 公式語法
    # =========================================================
    cat = Category.SYNTAX

    a1 = str(ws["A1"].value or "")
    suite.assert_true("Banner 含「跨期間趨勢分析」", cat,
                      "跨期間趨勢分析" in a1)

    a2 = str(ws["A2"].value or "")
    suite.assert_true("sub-banner 引用 Tbl歷史", cat,
                      "Tbl歷史" in a2)

    # 期間下拉
    suite.assert_true("B3 期間下拉預設值", cat,
                      ws[PERIOD_CELL].value in PERIODS)
    dvs = list(ws.data_validations.dataValidation)
    has_period_dv = False
    for dv in dvs:
        if all(p in str(dv.formula1 or "") for p in PERIODS):
            for r in dv.sqref.ranges:
                if "B3" in str(r):
                    has_period_dv = True
    suite.assert_true("B3 有期間 DataValidation", cat, has_period_dv)

    # Insight headline (C3)
    c3 = str(ws["C3"].value or "")
    suite.assert_true("Insight 引用 B3 期間", cat, "B3" in c3, detail=c3[:80])

    # 主表 header
    headers = ["KPI", "近 N 月走勢", "當月", "3 月前", "6 月前",
               "vs 去年同期", "趨勢"]
    for i, h in enumerate(headers, start=1):
        col = chr(ord("A") + i - 1)
        suite.assert_eq(f"主表 header {col}7", cat, ws[f"{col}7"].value, h)

    # 5 KPI 主表
    expected_kpis = [(8, "全般破獲率"), (9, "毒調率"), (10, "未到驗"),
                     (11, "交通達標率"), (12, "未破案件")]
    for row, name in expected_kpis:
        v = str(ws[f"A{row}"].value or "")
        suite.assert_true(f"A{row} 含 KPI 名「{name}」", cat, name in v,
                          detail=v[:40])

    # 5 KPI 各欄公式存在
    for row in range(8, 13):
        for col in ["B", "C", "D", "E", "F", "G"]:
            f = str(ws[f"{col}{row}"].value or "")
            suite.assert_true(f"主表 {col}{row} 公式存在", cat,
                              f.startswith("="), detail=f[:60])

    # B 欄 sparkline 引用 AR helper
    f = str(ws["B8"].value or "")
    suite.assert_true("Sparkline 引用 AR helper", cat,
                      "AR" in f, detail=f[:60])

    # C 欄當月用 INDEX(Tbl歷史)
    f = str(ws["C8"].value or "")
    suite.assert_true("當月用 INDEX Tbl歷史", cat,
                      "INDEX(Tbl歷史" in f, detail=f[:60])

    # F 欄 vs 去年同期 含資料不足 fallback
    f = str(ws["F8"].value or "")
    suite.assert_true("vs 去年同期 含「資料不足」fallback", cat,
                      "資料不足" in f, detail=f[:80])

    # G 欄趨勢 含 ↑/↓/→
    f = str(ws["G8"].value or "")
    for arrow in ["↑", "↓", "→"]:
        suite.assert_true(f"趨勢欄含 '{arrow}'", cat, arrow in f, detail=f[:80])

    # 反向 KPI（未到驗 = reverse=True）
    f_g10 = str(ws["G10"].value or "")
    suite.assert_true("反向 KPI（未到驗）顯示「改善 / 惡化」", cat,
                      "改善" in f_g10 and "惡化" in f_g10, detail=f_g10[:80])

    # Top 3 變化卡
    for row in [15, 16, 17]:
        f = str(ws[f"A{row}"].value or "")
        suite.assert_true(f"Top 3 row {row} label 存在", cat,
                          f != "" and ("最大" in f or "需關注" in f or "持續" in f),
                          detail=f[:40])
        c_fml = str(ws[f"C{row}"].value or "")
        suite.assert_true(f"Top 3 row {row} 變化公式", cat,
                          c_fml.startswith("="), detail=c_fml[:60])

    # Helper area
    suite.assert_true("Helper A100 KPI 名存在", cat,
                      "全般破獲率" in str(ws["A100"].value or ""))
    t100 = str(ws["T100"].value or "")
    suite.assert_true("Helper T100 取 Tbl歷史 資料", cat,
                      "INDEX(Tbl歷史" in t100, detail=t100[:60])

    # block 字元公式 (AF100)
    af100 = str(ws["AF100"].value or "")
    suite.assert_true("AF100 block 字元公式（CHOOSE）", cat,
                      "CHOOSE" in af100 and "▁" in af100, detail=af100[:80])

    # CONCAT (AR100)
    ar100 = str(ws["AR100"].value or "")
    suite.assert_true("AR100 CONCAT 12 個 block 字元", cat,
                      "CONCAT" in ar100, detail=ar100[:60])

    # =========================================================
    # 2. 邊界資料
    # =========================================================
    cat = Category.BOUNDARY

    # 歷史資料 < 12 月 處理（MAX(1, ...) 在 helper, INDEX fallback "")
    for cell in ["T100", "C8", "F8"]:
        f = str(ws[cell].value or "")
        suite.assert_true(f"{cell} 處理歷史資料不足", cat,
                          "IFERROR" in f, detail=f[:80])

    # 反向 KPI（未到驗、未破案件）— trend 判斷正確
    for row in [10, 12]:
        f = str(ws[f"G{row}"].value or "")
        suite.assert_true(f"反向 KPI row {row} 趨勢判斷", cat,
                          "改善" in f or "惡化" in f, detail=f[:80])

    # MAX(1, ...) 處理空表
    f = str(ws["C8"].value or "")
    suite.assert_true("空表處理 MAX(1,...)", cat, "MAX(1," in f, detail=f[:80])

    # =========================================================
    # 3. 性能基準
    # =========================================================
    cat = Category.PERFORMANCE

    fcount = 0
    for row in ws.iter_rows():
        for cell in row:
            v = cell.value
            if v and isinstance(v, str) and v.startswith("="):
                fcount += 1
    suite.assert_le("頁 12 公式總數 ≤ 250", cat, fcount, 250,
                    detail=f"actual={fcount}")

    # 揮發性
    volatile_counts = {"TODAY()": 0, "OFFSET(": 0, "INDIRECT(": 0}
    for row in ws.iter_rows():
        for cell in row:
            if cell.value and isinstance(cell.value, str):
                for k in volatile_counts:
                    volatile_counts[k] += cell.value.count(k)
    suite.assert_eq("不直接呼叫 TODAY()", cat, volatile_counts["TODAY()"], 0)
    suite.assert_le("OFFSET ≤ 3", cat, volatile_counts["OFFSET("], 3)
    suite.assert_le("INDIRECT ≤ 3", cat, volatile_counts["INDIRECT("], 3)

    # =========================================================
    # 4. 跨平台相容
    # =========================================================
    cat = Category.CROSS_PLATFORM

    has_filter = False
    for row in ws.iter_rows():
        for cell in row:
            if cell.value and isinstance(cell.value, str):
                if "FILTER(" in cell.value or "SORT(" in cell.value:
                    has_filter = True
    suite.assert_true("頁 12 無 FILTER/SORT 依賴", cat, not has_filter)

    # CONCAT
    has_concat = False
    for row in ws.iter_rows():
        for cell in row:
            if cell.value and isinstance(cell.value, str) and "CONCAT(" in cell.value:
                has_concat = True
                break
    suite.assert_true("使用 CONCAT（block 字元串接）", cat, has_concat)

    # 命名範圍
    wb_nrs = set(wb.defined_names)
    for nr in ["今日", "派出所名稱"]:
        suite.assert_true(f"命名範圍 '{nr}' 存在", cat, nr in wb_nrs)

    # 結構引用 Tbl歷史
    f = str(ws["C8"].value or "")
    suite.assert_true("使用 Tbl歷史 結構引用", cat, "Tbl歷史" in f)

    # =========================================================
    # 5. 真實情境
    # =========================================================
    cat = Category.REAL_WORLD

    # 凍結 None（使用者要求 2026-05-25 全頁無凍結）
    suite.assert_eq("凍結窗格 = None（使用者要求全頁無凍結）", cat,
                    ws.freeze_panes, None)

    # 5 大 KPI 全到位
    expected_kpi_names = ["全般破獲率", "毒調率", "未到驗", "交通達標率", "未破案件"]
    for row, name in zip(range(8, 13), expected_kpi_names):
        v = str(ws[f"A{row}"].value or "")
        suite.assert_true(f"KPI 「{name}」在 row {row}", cat,
                          name in v, detail=v[:30])

    # 期間 3 選項
    for period in PERIODS:
        suite.assert_true(f"期間「{period}」在 DataValidation", cat,
                          any(period in str(dv.formula1 or "") for dv in dvs))

    # 頁尾說明 Sparkline 跨平台
    f = str(ws["A19"].value or "")
    suite.assert_true("頁尾說明 Sparkline 跨平台", cat,
                      "Sparkline" in f or "Unicode" in f, detail=f[:80])

    # Section title 動態引用期間
    a5 = str(ws["A5"].value or "")
    suite.assert_true("Section title 動態引用 B3 期間", cat,
                      "B3" in a5, detail=a5[:80])

    # =========================================================
    # 6. 錯誤恢復
    # =========================================================
    cat = Category.ERROR_RECOVERY

    # 主表 5 KPI × 各欄全包 IFERROR
    for row in range(8, 13):
        for col in ["B", "C", "D", "E", "F", "G"]:
            f = str(ws[f"{col}{row}"].value or "")
            if col == "B":
                # B 是 sparkline 引用，間接 IFERROR 在 helper
                continue
            suite.assert_true(f"主表 {col}{row} IFERROR", cat,
                              "IFERROR" in f, detail=f[:60])

    # Helper 公式 IFERROR
    for cell in ["T100", "AR100"]:
        f = str(ws[cell].value or "")
        suite.assert_true(f"Helper {cell} IFERROR", cat, "IFERROR" in f)

    # block 字元公式 ISNUMBER 處理空值
    af100 = str(ws["AF100"].value or "")
    suite.assert_true("AF100 ISNUMBER 處理空值", cat,
                      "ISNUMBER" in af100, detail=af100[:80])

    # vs 去年同期：歷史不足時顯示「資料不足」
    f = str(ws["F8"].value or "")
    suite.assert_true("vs 去年同期 fallback「資料不足」", cat,
                      "資料不足" in f)

    return suite


if __name__ == "__main__":
    suite = run()
    print(suite.report())
    sys.exit(0 if suite.is_all_pass() else 1)
