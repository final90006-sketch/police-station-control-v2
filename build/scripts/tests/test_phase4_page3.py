"""test_phase4_page3.py — 頁 3 5 分鐘報告 6 大類測試"""
import sys
from pathlib import Path
sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

from openpyxl import load_workbook
from phases._base import OUTPUT_PATH
from tests.framework import TestSuite, Category


def run():
    suite = TestSuite("頁 3 長官 5 分鐘報告")
    wb = load_workbook(OUTPUT_PATH)
    ws = wb["長官5分鐘報告"]

    # === 1. 公式語法 ===
    cat = Category.SYNTAX
    a1 = str(ws["A1"].value or "")
    suite.assert_true("Banner 引用 派出所名稱", cat, "派出所名稱" in a1)
    suite.assert_true("Banner 含「勤務狀況報告」", cat, "勤務狀況報告" in a1)

    a3 = str(ws["A3"].value or "")
    suite.assert_true("Row 3 含「報告人」+「分局督導官」", cat,
                      "報告人" in a3 and "督導官" in a3)

    # 結論引用紅黃綠燈
    a6 = str(ws["A6"].value or "")
    suite.assert_true("結論引用 管制總覽!A25 紅燈", cat,
                      "管制總覽!A25" in a6)
    suite.assert_true("結論依紅/黃/綠分支", cat,
                      "紅燈" in a6 and "黃燈" in a6 and "綠燈" in a6)

    # 三大重點 sections
    for r, name in [(10, "刑案績效"), (14, "毒品調驗"), (18, "交通取締")]:
        v = str(ws[f"A{r}"].value or "")
        suite.assert_true(f"重點 row {r} 含「{name}」", cat, name in v,
                          detail=v[:40])

    # 各重點 3 個 bullets (rows 11-13, 15-17, 19-21) 都應公式
    for r in [11, 12, 13, 15, 16, 17, 19, 20, 21]:
        v = str(ws[f"A{r}"].value or "")
        suite.assert_true(f"bullet A{r} 是公式", cat, v.startswith("="),
                          detail=v[:50])

    # 行動方案區（淡黃輸入）
    a24 = str(ws["A24"].value or "")
    suite.assert_true("行動 1 placeholder 含「請所長手填」", cat,
                      "請所長手填" in a24, detail=a24[:60])

    # Top 3 員警引用頁 10
    for rank in range(1, 4):
        r = 30 + rank
        for col in ["A", "B", "C"]:
            v = str(ws[f"{col}{r}"].value or "")
            suite.assert_true(f"Top {rank} {col}{r} 引用 績效統計", cat,
                              "績效統計!" in v, detail=v[:40])

    # === 2. 邊界 ===
    cat = Category.BOUNDARY
    # 行動方案 4 行
    for r in range(24, 28):
        v = ws[f"A{r}"].value
        suite.assert_true(f"行動方案 row {r} 存在", cat, v is not None)

    # 結論 IFERROR 邏輯（COUNTIF + IF chain）
    a6 = str(ws["A6"].value or "")
    suite.assert_true("結論 IF chain 全綠燈 fallback", cat,
                      "全綠燈通過" in a6 or "持續維持" in a6)

    # === 3. 性能 ===
    cat = Category.PERFORMANCE
    fcount = sum(1 for row in ws.iter_rows() for c in row
                 if c.value and isinstance(c.value, str) and c.value.startswith("="))
    suite.assert_le("頁 3 公式總數 ≤ 80", cat, fcount, 80, detail=f"actual={fcount}")

    # 不直接 TODAY()
    has_today = any("TODAY()" in str(c.value or "")
                    for row in ws.iter_rows() for c in row)
    suite.assert_true("不直接呼叫 TODAY()", cat, not has_today)

    # === 4. 跨平台 ===
    cat = Category.CROSS_PLATFORM
    has_filter = any("FILTER(" in str(c.value or "") or "SORT(" in str(c.value or "")
                     for row in ws.iter_rows() for c in row)
    suite.assert_true("無 FILTER/SORT 依賴", cat, not has_filter)

    wb_nrs = set(wb.defined_names)
    for nr in ["派出所名稱", "今日"]:
        suite.assert_true(f"命名範圍 '{nr}' 存在", cat, nr in wb_nrs)

    # === 5. 真實情境 ===
    cat = Category.REAL_WORLD
    # 4 大區塊 sections
    sections = [(5, "結論"), (9, "三大重點"), (23, "下月行動方案"), (29, "Top 3 員警")]
    for r, name in sections:
        v = str(ws[f"A{r}"].value or "")
        suite.assert_true(f"Section row {r} 含「{name}」", cat, name in v,
                          detail=v[:40])

    # 簽章區（所長 + 督導官）
    a35 = str(ws["A35"].value or "")
    e35 = str(ws["E35"].value or "")
    suite.assert_true("簽章區所長", cat, "所長" in a35 and "簽章" in a35)
    suite.assert_true("簽章區督導官", cat, "督導官" in e35 and "簽章" in e35)

    # 凍結窗格 = None (使用者要求)
    suite.assert_eq("凍結窗格 = None", cat, ws.freeze_panes, None)

    # === 6. 錯誤恢復 ===
    cat = Category.ERROR_RECOVERY
    # 各 bullet IFERROR 包覆（含 IFERROR 或引用既有已 IFERROR 公式）
    for r in [11, 15, 16, 17]:
        v = str(ws[f"A{r}"].value or "")
        suite.assert_true(f"bullet A{r} 容錯（IFERROR 或 計算欄）", cat,
                          "IFERROR" in v or "Tbl" in v, detail=v[:60])

    # Top 3 IFERROR 容錯
    for col in ["A", "B", "C"]:
        v = str(ws[f"{col}31"].value or "")
        suite.assert_true(f"Top 1 {col}31 IFERROR 容錯", cat, "IFERROR" in v)

    return suite


if __name__ == "__main__":
    suite = run()
    print(suite.report())
    sys.exit(0 if suite.is_all_pass() else 1)
