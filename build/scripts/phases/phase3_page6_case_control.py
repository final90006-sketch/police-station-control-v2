"""
phase3_page6_case_control.py — Phase 3：頁 6 刑案管制（B 公文風）

v2.1 改：移除「視圖下拉」，改成 **兩區段並呈**（上尚未偵破 + 下已破獲未移送）
理由：派出所實務上 2 種狀態都要常看，分開呈現比下拉切換更直觀

版型：
  Row 1-2  Banner
  Row 3    spacer
  Row 4    維護指引
  Row 5    三階段流程圖（純視覺，已移送=灰色不細列）
  Row 6    spacer
  Row 7    KPI 縮影：紅(尚未偵破) + 黃(已破獲未移送) + 灰(已移送)
  Row 8    spacer
  Row 9    🔴 Section「尚未偵破清單」
  Row 10   主表 header (A-I, 9 欄)
  Row 11-25 尚未偵破 案件 15 列容量
  Row 26   spacer
  Row 27   🟡 Section「已破獲未移送清單」
  Row 28   主表 header
  Row 29-43 已破獲未移送 案件 15 列容量
  Row 44   spacer
  Row 45   頁尾

兩區段資料源（helper area rows 100-299，隱藏）：
  col R = 是否「尚未偵破」（boolean）
  col S = 尚未偵破 累計順序
  col T = 是否「已破獲未移送」
  col U = 已破獲未移送 累計順序
"""
import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

from phases._base import backup_existing, OUTPUT_PATH, get_logger, unmerge_all_in_sheet
import styles as S


SHEET = "刑案管制"
ROC_YEAR = "(YEAR(今日)-1911)"
CASE_TBL = "Tbl案件"

DISPLAY_ROWS = 15
HELPER_START = 100
HELPER_LEN = 200


def build_page6_case_control(wb, log):
    log.info("--- 頁 6 刑案管制（v2.1 兩區段並呈版）---")
    ws = wb[SHEET]

    cleared = unmerge_all_in_sheet(ws)
    log.info(f"  清空既有 merges：{cleared} 個")

    # === 欄寬 ===
    col_widths = {
        "A": 12, "B": 22, "C": 18, "D": 30, "E": 18,
        "F": 18, "G": 16, "H": 16, "I": 26,
    }
    for c, w in col_widths.items():
        ws.column_dimensions[c].width = w

    # === Row 1-2 Banner ===
    ws.merge_cells("A1:I1")
    S.set_cell(ws, "A1",
               "   🔍  刑案管制（兩區段並呈：尚未偵破 + 已破獲未移送）",
               font_key="banner", fill_key="banner",
               align_key="left", border_key="bottom_thick")
    ws.row_dimensions[1].height = S.ROW_HEIGHT["banner"]

    ws.merge_cells("A2:I2")
    sub_formula = (
        '="   "&警察局名稱&" · "&分局名稱&" · "&派出所名稱'
        '&" ｜ 中華民國 "&(YEAR(今日)-1911)&" 年累計（截至 "&TEXT(今日,"mm/dd")'
        '&"）｜ 資料源：Tbl案件 (SSOT)"'
    )
    S.set_cell(ws, "A2", sub_formula,
               font_key=S.font(11, color="white", italic=True),
               fill_key=S.fill("accent_dark"), align_key="left")
    ws.row_dimensions[2].height = S.ROW_HEIGHT["sub_banner"]

    # === Row 3 spacer ===
    ws.row_dimensions[3].height = 8

    # === Row 4 維護指引 ===
    ws.merge_cells("A4:I4")
    S.set_cell(ws, "A4",
               '="📌 派出所主管兩階段：尚未偵破（要破）+ 已破獲未移送（催偵查隊）｜ '
               '已移送屬地檢端不在此清單細列"',
               font_key=S.font(11, color="warn", italic=True),
               fill_key="warn_bg",
               align_key="center", border_key="all_thin")
    ws.row_dimensions[4].height = 28

    # === Row 5 三階段流程圖 ===
    stage_cells = [
        ("A5:B5", "● 尚未偵破",     "danger", "danger_bg"),
        ("C5:C5", " ➜ ",            "muted",  None),
        ("D5:E5", "● 已破獲未移送", "warn",   "warn_bg"),
        ("F5:F5", " ➜ ",            "muted",  None),
        ("G5:I5", "● 已移送（地檢）", "muted",   "excel_calc"),
    ]
    for rng, label, fg, bg in stage_cells:
        first_cell = rng.split(":")[0]
        ws.merge_cells(rng)
        S.set_cell(ws, first_cell, label,
                   font_key=S.font(14, bold=True, color=fg),
                   fill_key=S.fill(bg) if bg else None,
                   align_key="center",
                   border_key="all_thin")
    ws.row_dimensions[5].height = 44

    # === Row 6 spacer ===
    ws.row_dimensions[6].height = 8

    # === Row 7 KPI 縮影 3 卡 ===
    unsolved_count = (
        f'COUNTIFS({CASE_TBL}[案件狀況],"尚未偵破",'
        f'{CASE_TBL}[發生管轄],"本轄",'
        f'{CASE_TBL}[歸屬年度],{ROC_YEAR})'
    )
    not_sent_count = (
        f'COUNTIFS({CASE_TBL}[案件狀況],"已破獲未移送",'
        f'{CASE_TBL}[查獲管轄],"本轄",'
        f'{CASE_TBL}[歸屬年度],{ROC_YEAR})'
    )
    sent_count = (
        f'COUNTIFS({CASE_TBL}[案件狀況],"已移送",'
        f'{CASE_TBL}[查獲管轄],"本轄",'
        f'{CASE_TBL}[歸屬年度],{ROC_YEAR})'
    )
    kpi_cards = [
        ("A7:C7", f"=\"🔴 尚未偵破：\"&IFERROR({unsolved_count},0)&\" 件\"",
            "danger", "danger_bg"),
        ("D7:F7", f"=\"🟡 已破獲未移送：\"&IFERROR({not_sent_count},0)&\" 件\"",
            "warn", "warn_bg"),
        ("G7:I7", f"=\"✓ 已移送（地檢）：\"&IFERROR({sent_count},0)&\" 件\"",
            "muted", "excel_calc"),
    ]
    for rng, fml, fg, bg in kpi_cards:
        ws.merge_cells(rng)
        first = rng.split(":")[0]
        S.set_cell(ws, first, fml,
                   font_key=S.font(16, bold=True, color=fg),
                   fill_key=S.fill(bg),
                   align_key="center", border_key="all_thin")
    ws.row_dimensions[7].height = 40

    # === Row 8 spacer ===
    ws.row_dimensions[8].height = 12

    # =============================================================
    # 區段 1：🔴 尚未偵破清單 (rows 9-25)
    # =============================================================
    ws.merge_cells("A9:I9")
    S.set_cell(ws, "A9",
               '="   🔴  尚未偵破清單（"&IFERROR(' + unsolved_count + ',0)&" 件，最多顯示前 ' + str(DISPLAY_ROWS) + ' 件）"',
               font_key="section_title", fill_key=S.fill("danger"),
               align_key="left")
    ws.row_dimensions[9].height = S.ROW_HEIGHT["header"]

    # Row 10 header（v2.2 精簡：備註→自填案類）
    headers = ["編號", "案類", "發生時間", "發生地", "破獲時間",
               "案件狀況", "偵辦進度", "承辦人", "自填案類"]
    for i, h in enumerate(headers, start=1):
        col = get_column_letter(i)
        S.set_cell(ws, f"{col}10", h,
                   font_key="header", fill_key="header",
                   align_key="center", border_key="all_thin")
    ws.row_dimensions[10].height = S.ROW_HEIGHT["header"]

    # Row 11-25: 15 列 尚未偵破 INDEX/MATCH from helper col S
    case_columns = [
        ("A", "編號", "general"),
        ("B", "案類", "general"),
        ("C", "發生時間", "date_roc"),
        ("D", "發生地點", "general"),
        ("E", "破獲時間", "date_roc"),
        ("F", "案件狀況", "general"),
        ("G", "偵辦進度", "general"),
        ("H", "承辦人", "general"),
        ("I", "自填案類", "general"),
    ]
    helper_S_range = f"$S${HELPER_START}:$S${HELPER_START+HELPER_LEN-1}"
    helper_U_range = f"$U${HELPER_START}:$U${HELPER_START+HELPER_LEN-1}"

    for k in range(1, DISPLAY_ROWS + 1):
        row = 10 + k
        match_expr = f"MATCH({k},{helper_S_range},0)"
        for col, field, fmt in case_columns:
            fml = f'=IFERROR(INDEX({CASE_TBL}[{field}],{match_expr}),"")'
            S.set_cell(ws, f"{col}{row}", fml,
                       font_key="body", fill_key="calc",
                       align_key="left" if col in ("B", "D", "I") else "center",
                       border_key="all_thin",
                       number_format=fmt)
        ws.row_dimensions[row].height = S.ROW_HEIGHT["default"]

    # === Row 26 spacer ===
    ws.row_dimensions[26].height = 14

    # =============================================================
    # 區段 2：🟡 已破獲未移送清單 (rows 27-43)
    # =============================================================
    ws.merge_cells("A27:I27")
    S.set_cell(ws, "A27",
               '="   🟡  已破獲未移送清單（"&IFERROR(' + not_sent_count + ',0)&" 件，最多顯示前 ' + str(DISPLAY_ROWS) + ' 件，需催偵查隊）"',
               font_key="section_title", fill_key=S.fill("warn"),
               align_key="left")
    ws.row_dimensions[27].height = S.ROW_HEIGHT["header"]

    # Row 28 header (重複，方便閱讀)
    for i, h in enumerate(headers, start=1):
        col = get_column_letter(i)
        S.set_cell(ws, f"{col}28", h,
                   font_key="header", fill_key="header",
                   align_key="center", border_key="all_thin")
    ws.row_dimensions[28].height = S.ROW_HEIGHT["header"]

    # Row 29-43: 15 列 已破獲未移送 INDEX/MATCH from helper col U
    for k in range(1, DISPLAY_ROWS + 1):
        row = 28 + k
        match_expr = f"MATCH({k},{helper_U_range},0)"
        for col, field, fmt in case_columns:
            fml = f'=IFERROR(INDEX({CASE_TBL}[{field}],{match_expr}),"")'
            S.set_cell(ws, f"{col}{row}", fml,
                       font_key="body", fill_key="calc",
                       align_key="left" if col in ("B", "D", "I") else "center",
                       border_key="all_thin",
                       number_format=fmt)
        ws.row_dimensions[row].height = S.ROW_HEIGHT["default"]

    # === Row 44 spacer ===
    ws.row_dimensions[44].height = 12

    # === Row 45 頁尾 ===
    ws.merge_cells("A45:I45")
    footer = (
        '="© KKEVIN-LIN-2026-V2.0  ｜  v2.1 兩區段並呈（取代 v2 視圖下拉）  ｜  '
        '製表："&TEXT(今日,"e/mm/dd")'
    )
    S.set_cell(ws, "A45", footer,
               font_key=S.font(9, color="muted", italic=True),
               fill_key="calc", align_key="center")
    ws.row_dimensions[45].height = 22

    # =============================================================
    # Helper area (隱藏，rows 100-299)
    # =============================================================
    # Q 序號 / R 是否尚未偵破 / S 尚未偵破累計 / T 是否已破獲未移送 / U 已破獲未移送累計
    S.set_cell(ws, "Q99", "[helper]",
               font_key=S.font(9, color="muted"))
    S.set_cell(ws, "R99", "尚未偵破?",
               font_key=S.font(9, color="muted"))
    S.set_cell(ws, "S99", "尚未順序",
               font_key=S.font(9, color="muted"))
    S.set_cell(ws, "T99", "已破未送?",
               font_key=S.font(9, color="muted"))
    S.set_cell(ws, "U99", "已破順序",
               font_key=S.font(9, color="muted"))

    for k in range(1, HELPER_LEN + 1):
        helper_row = HELPER_START + k - 1
        # Q 序號
        S.set_cell(ws, f"Q{helper_row}", k,
                   font_key=S.font(9, color="muted"),
                   number_format="integer")
        # 對應 Tbl案件 row k 各欄
        status_at_k = f'IFERROR(INDEX({CASE_TBL}[案件狀況],{k}),"")'
        occur_at_k = f'IFERROR(INDEX({CASE_TBL}[發生管轄],{k}),"")'
        catch_at_k = f'IFERROR(INDEX({CASE_TBL}[查獲管轄],{k}),"")'
        year_at_k = f'IFERROR(INDEX({CASE_TBL}[歸屬年度],{k}),0)'

        # R 是否 尚未偵破
        r_fml = (
            f'=IFERROR(IF(AND({status_at_k}="尚未偵破",'
            f'{occur_at_k}="本轄",{year_at_k}={ROC_YEAR}),1,0),0)'
        )
        S.set_cell(ws, f"R{helper_row}", r_fml,
                   font_key=S.font(9, color="muted"),
                   number_format="integer")
        # S 尚未偵破 累計
        s_fml = f'=IF(R{helper_row}=1,SUM($R${HELPER_START}:R{helper_row}),0)'
        S.set_cell(ws, f"S{helper_row}", s_fml,
                   font_key=S.font(9, color="muted"),
                   number_format="integer")
        # T 是否 已破獲未移送
        t_fml = (
            f'=IFERROR(IF(AND({status_at_k}="已破獲未移送",'
            f'{catch_at_k}="本轄",{year_at_k}={ROC_YEAR}),1,0),0)'
        )
        S.set_cell(ws, f"T{helper_row}", t_fml,
                   font_key=S.font(9, color="muted"),
                   number_format="integer")
        # U 已破獲未移送 累計
        u_fml = f'=IF(T{helper_row}=1,SUM($T${HELPER_START}:T{helper_row}),0)'
        S.set_cell(ws, f"U{helper_row}", u_fml,
                   font_key=S.font(9, color="muted"),
                   number_format="integer")

    # 隱藏 helper
    for r in range(99, HELPER_START + HELPER_LEN):
        ws.row_dimensions[r].hidden = True
    for col in ["Q", "R", "S", "T", "U"]:
        ws.column_dimensions[col].hidden = True

    # 公式計數
    fcount = 0
    for row in ws.iter_rows():
        for cell in row:
            if cell.value and isinstance(cell.value, str) and cell.value.startswith("="):
                fcount += 1
    log.info(f"  v2.1 兩區段並呈：尚未偵破 {DISPLAY_ROWS} 列 + 已破獲未移送 {DISPLAY_ROWS} 列")
    log.info(f"  Helper：{HELPER_LEN} 列 × 4 公式（rows 99-299 隱藏）")
    log.info(f"  公式總數：{fcount}")
    return fcount


def build():
    log = get_logger("phase3_p6")
    log.info("===== Phase 3 — 頁 6 刑案管制 開始 =====")

    wb = load_workbook(OUTPUT_PATH)
    fcount = build_page6_case_control(wb, log)

    bak = backup_existing()
    if bak:
        log.info(f"  舊版備份：{bak}")
    wb.save(OUTPUT_PATH)
    size_kb = OUTPUT_PATH.stat().st_size / 1024
    log.info(f"  儲存：{OUTPUT_PATH.name} ({size_kb:.1f} KB)")
    log.info(f"  公式：{fcount}")
    log.info("===== Phase 3 — 頁 6 完成 =====\n")
    return OUTPUT_PATH


if __name__ == "__main__":
    build()
