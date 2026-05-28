"""
phase2_data_tables.py — Phase 2：三大資料表（案件/取締/歷史）

職責：
- 頁 5 案件資料庫（22 欄，Tbl案件，1000 列）
- 頁 9.5 交通取締明細（6 欄，Tbl交通，196 列）
- 頁 14 歷史資料（14 欄，Tbl歷史，140 列）

每張表：
- 覆寫 Phase 1 通用 banner，加成品 banner + 容量條 + 維護指引
- 表頭（凍結窗格列）
- 示範資料列（input 淡黃；calc 淡灰 + 公式）
- 預留空白列（淡黃，使用者直接 key 即可）
- Excel Table 包覆全範圍（自動擴展、命名欄）
- Data Validation 下拉清單（連動設定表）
"""
import sys
from datetime import datetime
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

from openpyxl import load_workbook
from openpyxl.worksheet.table import Table, TableStyleInfo, TableColumn, TableFormula
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.utils import get_column_letter

from phases._base import load_config, backup_existing, OUTPUT_PATH, get_logger
import styles as S


import re as _re


def _structured_to_plain(formula: str, row: int, columns: list) -> str:
    """把計算欄公式的結構引用 [@欄名] 轉成純儲存格參照 {欄字母}{row}。

    ★★ 防 #REF! 地雷的核心（v2.2 第二版修）：
    Excel Table 計算欄若用 [@欄名] 結構引用，使用者刪列/編輯時 Excel
    會把它崩成 #REF!（且補 calculatedColumnFormula 又會被 Excel 判定
    table 不合法而整個移除）。改用純參照 D15/B15，刪列自動位移、永不
    崩 #REF!，且其他頁的 Tbl案件[案類分類] 結構引用照常運作。

    例：[@發生時間] 在 row 15 → D15（發生時間 是第 4 欄）
    """
    name_to_col = {}
    for i, col in enumerate(columns, start=1):
        name_to_col[col["name"]] = get_column_letter(i)

    def repl(m):
        name = m.group(1)
        col_letter = name_to_col.get(name)
        return f"{col_letter}{row}" if col_letter else m.group(0)

    return _re.sub(r"\[@([^\]]+)\]", repl, formula)


def _coerce_value(val, fmt):
    """日期/datetime 欄位若收到字串，嘗試轉成 datetime 物件，否則 YEAR/MONTH 等 Excel 公式會 #VALUE!"""
    if not isinstance(val, str) or not val:
        return val
    if fmt not in ("datetime", "date_roc", "date_iso"):
        return val
    for f in ("%Y-%m-%d %H:%M", "%Y-%m-%d", "%Y/%m/%d %H:%M", "%Y/%m/%d"):
        try:
            return datetime.strptime(val, f)
        except ValueError:
            continue
    return val


def _override_banner(ws, title: str, capacity_info: str):
    """覆寫 Phase 1 banner，加上資料表專屬資訊"""
    S.set_cell(ws, "A1", f"   📋  {title}",
               font_key="banner",
               fill_key="banner",
               align_key="left",
               border_key="bottom_thick")
    S.set_cell(ws, "A2", f"   {capacity_info}",
               font_key=S.font(11, color="white", italic=True),
               fill_key=S.fill("accent_dark"),
               align_key="left")


def _build_one_table(ws, wb, tbl_cfg: dict, log):
    """通用建表函式：依 config 建一張資料表"""
    sheet_name = tbl_cfg["sheet"]
    table_name = tbl_cfg["table_name"]
    header_row = tbl_cfg["header_row"]
    data_start = tbl_cfg["data_start_row"]
    columns = tbl_cfg["columns"]
    samples = tbl_cfg.get("sample_rows", [])
    blank_rows = tbl_cfg.get("blank_rows", 30)
    capacity = tbl_cfg.get("capacity", 100)
    table_style = tbl_cfg.get("table_style", "TableStyleMedium2")

    log.info(f"--- {sheet_name} → {table_name} ---")

    # === 覆寫 banner ===
    cap_info = f"📊 容量：{len(samples)} 已用 / {capacity} 上限"
    if tbl_cfg.get("warning_threshold"):
        cap_info += f"  ｜  ⚠ 達 {tbl_cfg['warning_threshold']} 列系統檢核會亮黃燈"
    if tbl_cfg.get("password"):
        cap_info += f"  ｜  🔒 受保護（密碼：{tbl_cfg['password']}）"
    _override_banner(ws, sheet_name, cap_info)

    # === Row 3-12 維護指引 + 容量條 + 欄位設計縮影 ===
    # Row 3 spacer
    ws.row_dimensions[3].height = 8

    # Row 4-5 維護指引
    ws.merge_cells(start_row=4, start_column=1,
                    end_row=4, end_column=min(8, len(columns)))
    S.set_cell(ws, "A4",
               "   💡  新增資料：① 點下方任一空白列 → ② 直接 key 入。 Excel Table 自動擴展。"
               "淡黃色 = 可輸入欄位，淡灰色 = 公式自動算（請勿改）",
               font_key=S.font(11, color="warn", italic=True),
               fill_key=S.fill("warn_bg"),
               align_key="left",
               border_key="all_thin")
    ws.row_dimensions[4].height = 30

    # Row 5 SSOT 連動說明
    ws.merge_cells(start_row=5, start_column=1,
                    end_row=5, end_column=min(8, len(columns)))
    S.set_cell(ws, "A5",
               f"   🔗  SSOT：本表 {table_name} 是其他頁面公式的資料源。"
               "修改後其他頁的 KPI、報告、檢核都會即時更新",
               font_key=S.font(10, color="muted", italic=True),
               fill_key=S.fill("accent_light"),
               align_key="left",
               border_key="all_thin")
    ws.row_dimensions[5].height = 26

    # Rows 6-13 spacer for banner to header gap (freeze pane needs row 14 visible after freeze)
    for r in range(6, header_row):
        ws.row_dimensions[r].height = 6 if r != header_row - 1 else 16

    # === 表頭 row at header_row ===
    # 注意：不加 ★ 前綴！結構引用 [@案件狀況] 需要欄位名字完全一致，
    # 若 header 是 "★ 案件狀況"，[@案件狀況] 會被 Excel 自動「修復」成
    # 亂七八糟的 [#標題] 引用，calc 公式整個爛掉。
    # 改用 sub-header 一行示意「★ = 下拉欄位」。
    for i, col in enumerate(columns, start=1):
        col_letter = get_column_letter(i)
        header_text = col["name"]
        S.set_cell(ws, f"{col_letter}{header_row}", header_text,
                   font_key="header",
                   fill_key="header",
                   align_key="center",
                   border_key="all_thin")
        ws.column_dimensions[col_letter].width = col.get("width", 14)
    ws.row_dimensions[header_row].height = S.ROW_HEIGHT["header"]

    # === 示範資料列 ===
    cur_row = data_start
    for sample in samples:
        for i, col in enumerate(columns, start=1):
            col_letter = get_column_letter(i)
            val = sample[i - 1] if i - 1 < len(sample) else ""
            fill_key = "calc" if col["kind"] == "calc" else "input"

            # calc 欄：套公式而非值（v2.2 改純參照，防 #REF!）
            if col["kind"] == "calc" and col.get("formula"):
                plain_fml = _structured_to_plain(col["formula"], cur_row, columns)
                S.set_cell(ws, f"{col_letter}{cur_row}", plain_fml,
                           font_key="body",
                           fill_key="calc",
                           align_key="center",
                           border_key="all_thin",
                           number_format=col.get("format", "general"))
            else:
                # 日期字串→datetime 物件（避免 YEAR/MONTH 計算欄 #VALUE!）
                coerced = _coerce_value(val, col.get("format", "general"))
                S.set_cell(ws, f"{col_letter}{cur_row}", coerced,
                           font_key="body",
                           fill_key=fill_key,
                           align_key="center",
                           border_key="all_thin",
                           number_format=col.get("format", "general"))
        ws.row_dimensions[cur_row].height = S.ROW_HEIGHT["default"]
        cur_row += 1

    # === 預留空白列（含 calc 欄公式預填）===
    end_row = data_start + len(samples) + blank_rows - 1
    for r in range(cur_row, end_row + 1):
        for i, col in enumerate(columns, start=1):
            col_letter = get_column_letter(i)
            if col["kind"] == "calc" and col.get("formula"):
                plain_fml = _structured_to_plain(col["formula"], r, columns)
                S.set_cell(ws, f"{col_letter}{r}", plain_fml,
                           font_key="body",
                           fill_key="calc",
                           align_key="center",
                           border_key="all_thin",
                           number_format=col.get("format", "general"))
            else:
                cell = ws.cell(row=r, column=i)
                cell.fill = S.FILL["input"]
                cell.border = S.BORDER["all_thin"]
                if col.get("format") and col["format"] != "general":
                    cell.number_format = S.NUM_FMT.get(col["format"], col["format"])
        ws.row_dimensions[r].height = S.ROW_HEIGHT["default"]
    log.info(f"  示範 {len(samples)} 列 + 預留 {blank_rows} 列 = 表共 {end_row - header_row + 1} 列")

    # === Excel Table ===
    # ★ v2.2 第二版修：計算欄改純儲存格參照（見 _structured_to_plain），
    #   table 維持 openpyxl 自動產生欄定義（不加 calculatedColumnFormula
    #   —— 那會被 Excel 判定不合法而整個移除 table）。
    last_col = get_column_letter(len(columns))
    ref = f"A{header_row}:{last_col}{end_row}"
    table = Table(displayName=table_name, ref=ref)
    table.tableStyleInfo = TableStyleInfo(
        name=table_style,
        showRowStripes=True,
        showColumnStripes=False,
        showFirstColumn=False,
        showLastColumn=False)
    ws.add_table(table)
    calc_n = sum(1 for c in columns if c["kind"] == "calc" and c.get("formula"))
    log.info(f"  Excel Table：{table_name} 範圍 {ref} 樣式 {table_style}"
             f"（{calc_n} 計算欄用純參照防 #REF!）")

    # === Data Validation：下拉清單 ===
    for i, col in enumerate(columns, start=1):
        if col.get("dv"):
            col_letter = get_column_letter(i)
            dv = DataValidation(
                type="list",
                formula1=f"={col['dv']}",
                allow_blank=True,
                showErrorMessage=True,
                errorTitle="輸入值不合法",
                error=f"請從下拉清單選擇有效的「{col['name']}」"
            )
            dv.add(f"{col_letter}{data_start}:{col_letter}{end_row}")
            ws.add_data_validation(dv)
            log.info(f"  下拉：{col['name']} → {col['dv']}")

    # === Data Validation：日期欄（datetime/date_roc/date_iso 強制有效日期）===
    for i, col in enumerate(columns, start=1):
        fmt = col.get("format", "")
        if fmt in ("datetime", "date_roc", "date_iso") and col["kind"] == "input":
            col_letter = get_column_letter(i)
            dv_date = DataValidation(
                type="date",
                operator="between",
                formula1="DATE(2020,1,1)",
                formula2="DATE(2035,12,31)",
                allow_blank=True,
                showErrorMessage=True,
                errorTitle="日期格式錯誤",
                error=(f"「{col['name']}」必須是有效日期\n"
                       "✓ 範例：2026/05/26 或 2026/05/26 14:00\n"
                       "✗ 不可輸入文字、2020 以前、2035 以後"),
                showInputMessage=True,
                promptTitle=f"{col['name']} — 日期格式",
                prompt=("請輸入有效日期（顯示精度到「時」）\n"
                        "範例：2026/05/26（無時間）\n"
                        "或：2026/05/26 14:00（含小時，分鐘會被隱藏）")
            )
            dv_date.add(f"{col_letter}{data_start}:{col_letter}{end_row}")
            ws.add_data_validation(dv_date)
            log.info(f"  日期驗證：{col['name']} → DATE 範圍 2020-2035")

    return end_row


def _add_case_summary(ws, log):
    """案件資料庫頂端即時統計（全般 / 竊盜 / 詐欺 × 發生 / 破獲 / 破獲率）。

    使用者建議：列管刑案時隨時掌握數字，不必跳到頁 2。
    放 rows 6-10（_build_one_table 預留的 spacer 區），本轄當年。
    """
    ROC = "(YEAR(今日)-1911)"
    T = "Tbl案件"

    def occ(cat=None):
        base = f'COUNTIFS({T}[發生管轄],"本轄",{T}[歸屬年度],{ROC},{T}[計入發生數],"是"'
        if cat:
            base = f'COUNTIFS({T}[案類分類],"{cat}",{T}[發生管轄],"本轄",{T}[歸屬年度],{ROC},{T}[計入發生數],"是"'
        return base + ")"

    def solved(cat=None):
        base = f'COUNTIFS({T}[查獲管轄],"本轄",{T}[是否破獲],"是",{T}[歸屬年度],{ROC}'
        if cat:
            base = f'COUNTIFS({T}[案類分類],"{cat}",{T}[查獲管轄],"本轄",{T}[是否破獲],"是",{T}[歸屬年度],{ROC}'
        return base + ")"

    # Row 6 title
    ws.merge_cells("A6:I6")
    S.set_cell(ws, "A6",
               '="   📊  即時管制統計（中華民國 "&(YEAR(今日)-1911)&" 年・本轄）'
               ' ｜ 隨資料即時更新，列管刑案隨時掌握數字"',
               font_key="section_title", fill_key="banner", align_key="left")
    ws.row_dimensions[6].height = 28

    # Row 7 header
    sum_headers = [("A", "案類"), ("B", "發生"), ("C", "破獲"), ("D", "破獲率")]
    for col, h in sum_headers:
        S.set_cell(ws, f"{col}7", h, font_key="header", fill_key="header",
                   align_key="center", border_key="all_thin")
    ws.row_dimensions[7].height = 24

    # Row 8-10 全般 / 竊盜 / 詐欺
    cats = [("全般", None, "accent"), ("竊盜", "竊盜", "text"), ("詐欺", "詐欺", "text")]
    for i, (label, cat, color) in enumerate(cats):
        r = 8 + i
        S.set_cell(ws, f"A{r}", label, font_key=S.font(12, bold=True, color=color),
                   fill_key="accent_light" if cat is None else "calc",
                   align_key="center", border_key="all_thin")
        S.set_cell(ws, f"B{r}", f"={occ(cat)}",
                   font_key=S.font(14, bold=True, color="accent"),
                   fill_key="calc", align_key="center", border_key="all_thin",
                   number_format="integer")
        S.set_cell(ws, f"C{r}", f"={solved(cat)}",
                   font_key=S.font(14, bold=True, color="pass"),
                   fill_key="calc", align_key="center", border_key="all_thin",
                   number_format="integer")
        S.set_cell(ws, f"D{r}", f"=IFERROR(C{r}/B{r},0)",
                   font_key=S.font(14, bold=True, color="warn"),
                   fill_key="calc", align_key="center", border_key="all_thin",
                   number_format="percent_one")
        ws.row_dimensions[r].height = 24
    log.info("  案件資料庫：頂端即時統計（全般/竊盜/詐欺 × 發生/破獲/破獲率）")


def build():
    log = get_logger("phase2")
    log.info("===== Phase 2（三大資料表）開始 =====")

    cfg = load_config("phase2_data_tables")
    wb = load_workbook(OUTPUT_PATH)

    # === 頁 5 案件資料庫 ===
    p5 = cfg["page5_case_database"]
    ws5 = wb[p5["sheet"]]
    _build_one_table(ws5, wb, p5, log)
    _add_case_summary(ws5, log)   # v2.2 使用者建議：頂端即時統計

    # === 頁 9.5 交通取締明細 ===
    p95 = cfg["page9_5_traffic_detail"]
    ws95 = wb[p95["sheet"]]
    _build_one_table(ws95, wb, p95, log)

    # === 頁 14 歷史資料 ===
    p14 = cfg["page14_history"]
    ws14 = wb[p14["sheet"]]
    _build_one_table(ws14, wb, p14, log)

    # === 儲存 ===
    bak = backup_existing()
    if bak:
        log.info(f"  舊版備份：{bak}")
    wb.save(OUTPUT_PATH)
    size_kb = OUTPUT_PATH.stat().st_size / 1024
    log.info(f"  儲存：{OUTPUT_PATH.name} ({size_kb:.1f} KB)")
    log.info(f"  總命名範圍數：{len(wb.defined_names)}")
    log.info("===== Phase 2 完成 =====\n")
    return OUTPUT_PATH


if __name__ == "__main__":
    build()
