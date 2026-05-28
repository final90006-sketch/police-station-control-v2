"""
make_project_control.py — 產生「派出所綜合管制系統 v2.2 專案管制表」

輸出到桌面，記錄整個專案：開發階段、上線問題修復、交付物、待辦。
官方藍風格，A4 直式可列印。
"""
import sys
from pathlib import Path
from datetime import datetime

sys.path.insert(0, str(Path(__file__).resolve().parent))

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

OUT = r"C:\Users\User\Desktop\派出所系統_專案管制表_v2.2.xlsx"

# 顏色
ACCENT = "FF1F4E79"
ACCENT_DARK = "FF14253D"
ACCENT_LIGHT = "FFDBEAFE"
GREEN = "FF15803D"; GREEN_BG = "FFDCFCE7"
WARN = "FFB45309"; WARN_BG = "FFFEF3C7"
GREY = "FF64748B"; CALC_BG = "FFF1F5F9"
WHITE = "FFFFFFFF"

thin = Side("thin", color="FFE2E8F0")
BORDER = Border(left=thin, right=thin, top=thin, bottom=thin)
CTR = Alignment(horizontal="center", vertical="center", wrap_text=True)
LEFT = Alignment(horizontal="left", vertical="center", wrap_text=True)


def cell(ws, coord, val, *, size=11, bold=False, color="FF1E293B",
         bg=None, align=CTR, border=True):
    c = ws[coord]
    c.value = val
    c.font = Font(name="微軟正黑體", size=size, bold=bold, color=color)
    if bg:
        c.fill = PatternFill("solid", fgColor=bg)
    c.alignment = align
    if border:
        c.border = BORDER
    return c


def main():
    wb = Workbook()
    ws = wb.active
    ws.title = "專案管制表"
    ws.sheet_view.showGridLines = False

    widths = {"A": 6, "B": 16, "C": 40, "D": 12, "E": 14, "F": 30}
    for col, w in widths.items():
        ws.column_dimensions[col].width = w

    # === Banner ===
    ws.merge_cells("A1:F1")
    cell(ws, "A1", "  📋  派出所綜合管制系統 v2.2 — 專案管制表",
         size=18, bold=True, color=WHITE, bg=ACCENT, align=LEFT, border=False)
    ws.row_dimensions[1].height = 44

    ws.merge_cells("A2:F2")
    cell(ws, "A2",
         f"  製表：林錦瑞 ｜ KKEVIN-LIN-2026-V2.2 ｜ 更新：{datetime.now():%Y/%m/%d} ｜ "
         "GitHub: github.com/final90006-sketch/police-station-control-v2",
         size=10, color=WHITE, bg=ACCENT_DARK, align=LEFT, border=False)
    ws.row_dimensions[2].height = 24

    # === 摘要列 ===
    ws.row_dimensions[3].height = 6
    summary = [
        ("A4:B4", "整體完成度", "C4:C4", "工程 100% ｜ 部署準備 100% ｜ 待實機試用", GREEN, GREEN_BG),
        ("A5:B5", "公式測試", "C5:C5", "1426 / 1426 全綠燈（14 測試檔・6 大類框架）", GREEN, GREEN_BG),
        ("A6:B6", "系統規模", "C6:C6", "17 分頁 ｜ 4 大資料表 ｜ 45 條系統檢核 ｜ 約 138 KB", ACCENT, ACCENT_LIGHT),
    ]
    for lab_rng, lab, val_rng, val, fg, bg in summary:
        ws.merge_cells(lab_rng)
        cell(ws, lab_rng.split(":")[0], lab, bold=True, color=fg, bg=bg, align=LEFT)
        ws.merge_cells(val_rng.replace("C", "C").split(":")[0] + ":F" + lab_rng[1])
        cell(ws, val_rng.split(":")[0], val, color="FF1E293B", bg=bg, align=LEFT)
    for r in (4, 5, 6):
        ws.row_dimensions[r].height = 22

    # === 表頭 ===
    ws.row_dimensions[7].height = 8
    headers = ["項次", "階段", "工作項目", "狀態", "完成日", "交付物 / 備註"]
    for i, h in enumerate(headers, start=1):
        cell(ws, f"{get_column_letter(i)}8", h, size=12, bold=True,
             color=WHITE, bg=ACCENT)
    ws.row_dimensions[8].height = 30

    # === 管制項目 ===
    # (階段, 項目, 狀態, 完成日, 交付物/備註)
    DONE = ("✓ 完成", GREEN, GREEN_BG)
    PEND = ("⏳ 待辦", WARN, WARN_BG)
    rows = [
        ("壹 系統開發", "Phase 1-2：17 工作表骨架 + 設定表 8 區 + 4 大資料表", *DONE, "0526", "Tbl案件/交通/歷史/毒調"),
        ("壹 系統開發", "Phase 3：14 頁實質內容（九宮格/環形圖/兩區段/毒調）", *DONE, "0526", "頁 1-12 全建置"),
        ("壹 系統開發", "Phase 4：長官 5 分鐘報告 + 對上級上呈", *DONE, "0526", "Pyramid / CHOOSE 三格式"),
        ("壹 系統開發", "Phase 5：條件格式 + 工作表保護 KD2026 + 列印區", *DONE, "0526", "紅黃綠燈 / 11 頁上鎖"),
        ("壹 系統開發", "Phase 6：頁 15 系統檢核 45 條（7 大類）", *DONE, "0526", "含跳轉連結"),
        ("壹 系統開發", "Phase 7：Tbl毒調接通 + 頁 2 KPI 4/5", *DONE, "0526", "毒調率 / 未到驗"),
        ("壹 系統開發", "Phase 8：頁 14 即時 KPI 抄寫區（員警冠軍自動）", *DONE, "0526", "月底抄寫防呆"),
        ("壹 系統開發", "Phase 9：真原生 Excel Sparkline（XML 注入）", *DONE, "0526", "頁 2 七條折線"),
        ("壹 系統開發", "Phase 10：最後到驗日 30 日 / sparkline / 員警冠軍", *DONE, "0526", "ABC 強化"),
        ("貳 部署發布", "Phase 11：部署版產出（清 sample 資料）", *DONE, "0526", "PoliceStation_v2.2_部署版.xlsx"),
        ("貳 部署發布", "部署手冊 SOP（安裝 / 月例行 / 督導稽核）", *DONE, "0526", "README_部署手冊_v2.2.md"),
        ("貳 部署發布", "超白話操作說明 PDF（10 頁圖文）", *DONE, "0526", "派出所管制系統_超白話操作說明.pdf"),
        ("貳 部署發布", "Phase 12：GitHub 公開發布（CC BY 4.0）", *DONE, "0526", "police-station-control-v2"),
        ("參 上線修復", "🔴 #REF! 地雷根治（計算欄改純儲存格參照）", *DONE, "0528", "build + 同仁檔皆修"),
        ("參 上線修復", "同仁資料清洗：159 筆文字日期 → 真日期", *DONE, "0528", "含民國年 / HHMM / typo"),
        ("參 上線修復", "同仁資料清洗：85 筆案類 → 下拉字典值", *DONE, "0528", "原文存自填案類"),
        ("參 上線修復", "8 筆法院交辦：發生本轄 + 查獲依破獲時間判定", *DONE, "0528", "有破獲時間=我們破"),
        ("參 上線修復", "案件資料庫精簡 19 → 15 輸入欄", *DONE, "0528", "拿掉備註/金額/建立日期/建立者"),
        ("參 上線修復", "同仁檔遷移 15 欄 + 92 筆資料保留", *DONE, "0528", "PoliceStation_v2.2_最終版.xlsx"),
        ("肆 待辦（部署階段）", "跨平台實機測試（Win / Mac / LibreOffice）", *PEND, "—", "使用者端驗證"),
        ("肆 待辦（部署階段）", "列印實機測試（A4 直/橫，各廠牌印表機）", *PEND, "—", "頁 3/4/11/2.5"),
        ("肆 待辦（部署階段）", "三所試用 1 個月（小 / 中 / 大型派出所）", *PEND, "—", "回饋後微調"),
        ("肆 待辦（部署階段）", "其他同仁推廣（白話 PDF + 部署版）", *PEND, "—", "教育訓練"),
    ]

    r = 9
    for idx, (stage, item, status, fg, bg, date, deliver) in enumerate(rows, start=1):
        cell(ws, f"A{r}", idx, bg=CALC_BG)
        cell(ws, f"B{r}", stage, size=10, color=ACCENT, bg=CALC_BG, align=LEFT)
        cell(ws, f"C{r}", item, align=LEFT)
        cell(ws, f"D{r}", status, bold=True, color=fg, bg=bg)
        cell(ws, f"E{r}", date, color=GREY)
        cell(ws, f"F{r}", deliver, size=10, color=GREY, align=LEFT)
        ws.row_dimensions[r].height = 26
        r += 1

    # === 統計列 ===
    done_n = sum(1 for x in rows if x[2].startswith("✓"))
    pend_n = sum(1 for x in rows if x[2].startswith("⏳"))
    ws.merge_cells(f"A{r}:F{r}")
    cell(ws, f"A{r}",
         f"  共 {len(rows)} 項 ｜ ✓ 完成 {done_n} 項 ｜ ⏳ 待辦 {pend_n} 項"
         f"（皆屬部署實機階段，工程開發 100% 完成）",
         size=11, bold=True, color=WHITE, bg=ACCENT, align=LEFT)
    ws.row_dimensions[r].height = 28
    r += 1

    # === 頁尾 ===
    ws.merge_cells(f"A{r}:F{r}")
    cell(ws, f"A{r}",
         "  © 林錦瑞 ｜ 全台 1300+ 派出所通用 ｜ CC BY 4.0 ｜ 本表隨專案進度更新",
         size=9, color=GREY, align=LEFT, border=False)
    ws.row_dimensions[r].height = 20

    # 列印設定 A4 直式
    ws.print_area = f"A1:F{r}"
    ws.page_setup.orientation = "portrait"
    ws.page_setup.paperSize = 9
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 1
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_margins.left = ws.page_margins.right = 0.4
    ws.page_margins.top = ws.page_margins.bottom = 0.5

    wb.save(OUT)
    print(f"✓ 專案管制表：{OUT}")
    print(f"  共 {len(rows)} 項（完成 {done_n} / 待辦 {pend_n}）")


if __name__ == "__main__":
    main()
