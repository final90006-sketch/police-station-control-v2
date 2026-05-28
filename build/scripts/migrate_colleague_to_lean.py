"""
migrate_colleague_to_lean.py — 同仁檔遷移到 15 欄精簡版

把已清洗的同仁檔（25 欄）刪掉 4 個空欄（備註/涉案金額/建立日期/建立者），
變成 21 欄（15 input + 6 calc）的精簡結構，保留：
  - 設定表（含自訂員警名冊）
  - 92 筆案件資料
  - 所有其他頁

做法：
  1. 刪 Tbl案件 定義
  2. delete_cols(14, 4) 刪 N/O/P/Q
  3. 重建 Tbl案件 A14:U{end}
  4. 重寫 6 計算欄純參照（新位置）
"""
import sys
import os
import re
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent))

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter, column_index_from_string
from openpyxl.worksheet.table import Table, TableStyleInfo

from phases._base import load_config

SRC = r"D:\複本 PoliceStation_v2.2_最終清洗版.xlsx"
DST = r"D:\複本 PoliceStation_v2.2_精簡15欄_正式版.xlsx"


def structured_to_plain(formula, row, name_to_col):
    def repl(m):
        col = name_to_col.get(m.group(1))
        return f"{col}{row}" if col else m.group(0)
    return re.sub(r"\[@([^\]]+)\]", repl, formula)


def main():
    print("===== 同仁檔遷移 15 欄精簡版 開始 =====")
    if not os.path.exists(SRC):
        print(f"  ✗ 來源不存在：{SRC}")
        return

    wb = load_workbook(SRC)
    ws = wb["案件資料庫"]

    # 原 table 範圍
    old_ref = ws.tables["Tbl案件"].ref
    m = re.match(r"[A-Z]+(\d+):([A-Z]+)(\d+)", old_ref)
    header_row = int(m.group(1))
    end_row = int(m.group(3))
    print(f"  原 Tbl案件: {old_ref}（header row {header_row}, end {end_row}）")

    # 1. 刪 table 定義（避免 delete_cols 後 table 範圍錯亂）
    del ws.tables["Tbl案件"]

    # 2. 刪 N/O/P/Q（cols 14-17：備註/涉案金額/建立日期/建立者）
    #    確認這 4 欄 header 正確才刪
    to_del = [ws[f"{get_column_letter(c)}{header_row}"].value for c in range(14, 18)]
    print(f"  待刪 4 欄: {to_del}")
    expected = ["備註", "涉案金額", "建立日期", "建立者"]
    if to_del != expected:
        print(f"  ⚠ 欄位不符預期 {expected}，中止以策安全")
        return
    ws.delete_cols(14, 4)

    # 3. 重建 Tbl案件（新範圍 A14:U{end}）
    last_col = get_column_letter(21)   # 25-4=21 → U
    new_ref = f"A{header_row}:{last_col}{end_row}"
    tbl = Table(displayName="Tbl案件", name="Tbl案件", ref=new_ref)
    tbl.tableStyleInfo = TableStyleInfo(
        name="TableStyleMedium2", showFirstColumn=False, showLastColumn=False,
        showRowStripes=True, showColumnStripes=False)
    ws.add_table(tbl)
    print(f"  新 Tbl案件: {new_ref}")

    # 4. 重寫 6 計算欄純參照（新位置）
    cfg = load_config("phase2_data_tables")
    case_calc = {c["name"]: c["formula"].lstrip("=")
                 for c in cfg["page5_case_database"]["columns"]
                 if c["kind"] == "calc" and c.get("formula")}
    # 新 header → col
    name_to_col = {}
    for c in range(1, 22):
        cl = get_column_letter(c)
        h = ws[f"{cl}{header_row}"].value
        if h:
            name_to_col[h] = cl
    data_start = header_row + 1
    rewritten = 0
    for name, fstruct in case_calc.items():
        if name in name_to_col:
            col = name_to_col[name]
            for r in range(data_start, end_row + 1):
                ws[f"{col}{r}"] = "=" + structured_to_plain(fstruct, r, name_to_col)
                rewritten += 1
    print(f"  重寫 {rewritten} 計算欄 cell（純參照，新位置）")

    wb.calculation.fullCalcOnLoad = True
    wb.save(DST)

    # 驗證
    wb2 = load_workbook(DST)
    ws2 = wb2["案件資料庫"]
    print()
    print("=== 驗證 ===")
    print(f"  新 Tbl案件 範圍: {[t.ref for t in ws2.tables.values()]}")
    hdrs = [ws2[f'{get_column_letter(c)}{header_row}'].value for c in range(1, 22)]
    print(f"  21 欄 header: {hdrs}")
    print(f"  確認無 備註/涉案金額/建立日期/建立者: "
          f"{'✓' if not any(x in hdrs for x in expected) else '🔴'}")
    # 計算欄抽查
    n_col = name_to_col.get("歸屬年度")
    print(f"  歸屬年度({n_col}15): {str(ws2[f'{n_col}15'].value)[:45]}")
    print(f"  資料: row15 編號={ws2['A15'].value} / row106 編號={ws2['A106'].value}")
    print(f"  ✓ 遷移完成：{os.path.basename(DST)} ({os.path.getsize(DST)/1024:.1f} KB)")
    print("===== 完成 =====")


if __name__ == "__main__":
    main()
